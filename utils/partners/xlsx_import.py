# Übernimmt Partnerlinks, Status und Rabattcodes aus docs/partnerprogramme.xlsx nach db/partners.json.
# Aufruf: python3 utils/partners/xlsx_import.py [pfad/zur/datei.xlsx]
# Danach: npm run partners (prüft das Register und schreibt docs/PARTNERLINKS.md), dann bauen.
#
# Gelesen werden die Blätter „Tools“, „Verzeichnisse“, „KI-Tool-Datenbank“, „Neu entdeckt“ und „Alle Domains“.
# Eine Zeile wird übernommen, sobald Status, Partner-Link oder Rabattcode ausgefüllt bzw. vom Vorgabewert abweichend ist.
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / 'docs' / 'partnerprogramme.xlsx'
REG = ROOT / 'db' / 'partners.json'

register = json.loads(REG.read_text())
by_id = {p['id']: p for p in register['programme']}
STATUS = register['status_werte']


def leer(v):
    return v is None or str(v).strip() in ('', '–', '-')


def text(v):
    return '' if leer(v) else str(v).strip()


def zeilen(ws):
    kopf = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            yield dict(zip(kopf, row))


def eintrag(pid, name, kind, **felder):
    p = by_id.get(pid)
    if not p:
        p = {
            'id': pid,
            'name': name,
            'kind': kind,
            'tool_slug': None,
            'directory_slug': None,
            'network': '',
            'commission': '',
            'cookie': '',
            'signup_url': '',
            'status': 'nicht registriert',
            'partner_link': '',
            'coupon': None,
            'notes': '',
        }
        register['programme'].append(p)
        by_id[pid] = p
    for k, v in felder.items():
        if v is not None:
            p[k] = v
    return p


geaendert = []
erledigt = set()
wb = load_workbook(XLSX, data_only=True)

# Domain -> bestehender Eintrag, damit eine Domain aus „Alle Domains“ beim richtigen Tool oder Verzeichnis landet.
import subprocess
from urllib.parse import urlparse

host_von = lambda u: urlparse(u or '').netloc.lower().removeprefix('www.')
tool_host = {host_von(t['url']): t['slug'] for t in json.loads((ROOT / 'db/tools.json').read_text())}
dir_host = {
    host_von(d['link']): d['slug']
    for d in json.loads(subprocess.check_output(['node', '-e', "const s=require('./utils/slugify');const d=require('./db/directories.json');console.log(JSON.stringify((d.data||d).map(x=>({link:x.link,slug:s(x.name)}))))"], cwd=ROOT))
}


def finde(feld, key):
    for p in register['programme']:
        if p.get(feld) == key:
            return p
    if feld == 'host':
        if key in tool_host:
            return next((p for p in register['programme'] if p.get('tool_slug') == tool_host[key]), None)
        if key in dir_host:
            return next((p for p in register['programme'] if p.get('directory_slug') == dir_host[key]), None)
    return None

# Blattname -> (Spalte mit eindeutiger ID, Art, Zuordnungsfeld)
BLAETTER = {
    'Tools': ('Slug', 'tool', 'tool_slug'),
    'Verzeichnisse': ('Slug', 'directory', 'directory_slug'),
    'KI-Tool-Datenbank': ('Domain', 'kb-tool', 'host'),
    'Neu entdeckt': ('Domain', 'directory-neu', 'host'),
    'Alle Domains': ('Domain', 'domain', 'host'),
}

for blatt, (idspalte, kind, feld) in BLAETTER.items():
    if blatt not in wb.sheetnames:
        continue
    for z in zeilen(wb[blatt]):
        key = text(z.get(idspalte))
        if not key:
            continue
        status = text(z.get('Status'))
        link = text(z.get('Partner-Link (eintragen)'))
        code = text(z.get('Rabattcode'))
        rabatt = text(z.get('Rabatt-Text (englisch, z. B. 20% discount)'))
        notiz = text(z.get('Notizen'))
        bestehend = finde(feld, key)
        if bestehend and bestehend['id'] in erledigt:
            continue  # schon über ein vorheriges Blatt geändert (z. B. Tools vor „Alle Domains“)
        alt_status = bestehend['status'] if bestehend else 'nicht registriert'
        alt_code = ((bestehend or {}).get('coupon') or {}).get('code', '')
        # Nur übernehmen, was in der Tabelle neu oder anders ist als im Register (die Tabelle ist damit vorbefüllt).
        if not (
            (link and link != (bestehend or {}).get('partner_link'))
            or (code and code != alt_code)
            or (notiz and notiz != (bestehend or {}).get('notes'))
            or (status and status != alt_status)
        ):
            continue
        if status and status not in STATUS:
            print(f'Unbekannter Status in {blatt}, {key}: {status!r} (erlaubt: {", ".join(STATUS)})')
            continue
        if link and not re.match(r'^https://', link):
            print(f'Partner-Link muss mit https:// beginnen: {blatt}, {key}')
            continue
        pid = bestehend['id'] if bestehend else (key if feld != 'host' else re.sub(r'[^a-z0-9]+', '-', key.lower()).strip('-'))
        felder = {} if (bestehend and feld == 'host' and (bestehend.get('tool_slug') or bestehend.get('directory_slug'))) else {feld: key}
        if status:
            felder['status'] = status
        if link:
            felder['partner_link'] = link
            if status in ('', 'nicht registriert', 'beantragt', 'kein Programm'):
                felder['status'] = 'aktiv'
        if code:
            felder['coupon'] = {'code': code, 'text': rabatt or 'discount'}
        if notiz:
            felder['notes'] = notiz
        for spalte, ziel in (('Anmeldung (Registrierungslink)', 'signup_url'), ('Netzwerk', 'network'), ('Provision', 'commission'), ('Cookie', 'cookie')):
            wert = text(z.get(spalte))
            if wert and not (bestehend and bestehend.get(ziel)):
                felder[ziel] = wert
        vorher = json.dumps(by_id.get(pid), sort_keys=True)
        p = eintrag(pid, text(z.get('Name')) or key, kind, **felder)
        erledigt.add(p['id'])
        if json.dumps(p, sort_keys=True) != vorher:
            geaendert.append(f'{blatt}: {p["name"]} -> {p["status"]}{" mit Link" if p.get("partner_link") else ""}')

if geaendert:
    from datetime import date

    register['stand'] = date.today().isoformat()
    REG.write_text(json.dumps(register, ensure_ascii=False, indent=1) + '\n')
print(f'{len(geaendert)} Einträge geändert')
for g in geaendert:
    print(' ', g)
